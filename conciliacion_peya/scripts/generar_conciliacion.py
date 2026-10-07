#!/usr/bin/env python3
"""
Genera el Excel de conciliacion con columnas Concepto / Debe / Haber / Saldo,
mas las hojas de detalle.

El saldo arrastra con formula (=saldo_anterior + Debe - Haber), asi que se
puede tocar un importe y ver como se mueve el cierre sin rehacer nada.

USO
---
    python generar_conciliacion.py --json conceptos.json --salida Conciliacion.xlsx

FORMATO DEL JSON
----------------
{
  "titulo": "Conciliacion ...",
  "saldo_inicial": 83178161.80,
  "lineas": [
    {"concepto": "Retencion IIBB 16/03-22/03 Duplicada", "debe": 167988.04},
    {"concepto": "Ajuste Mal cargado Febrero", "haber": 6517.00},
    {"concepto": "Factura Pendiente 0026-02688968", "haber": 3528614.69, "flag": true},
    {"subtotal": "Subtotal"}
  ],
  "abierto": [{"concepto": "...", "debe": 63143.42, "nota": "..."}],
  "historico": "historico.csv",
  "detalles": {"Recaudacion": "mayor.xlsx", "Venta Dean": "dean.csv"}
}

"historico" es la hoja "Diferencias por mes": una fila por concepto de cruce,
una columna por periodo cerrado, mas el acumulado. Perdura entre cierres, y es
lo que permite que el detalle crudo del Excel sea solo del mes procesado.

Cada linea lleva "debe" o "haber", nunca las dos. "flag": true la pinta en
naranja (importe deducido o pendiente de definicion).
"""
from __future__ import annotations

import argparse
import datetime
import json
import os

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

F = "Arial"
FMT = "#,##0.00"
HEAD = PatternFill("solid", fgColor="1F4E78")
SALDO = PatternFill("solid", fgColor="FFF2CC")
FLAG = PatternFill("solid", fgColor="FCE4D6")
_T = Side(style="thin", color="BFBFBF")
BD = Border(left=_T, right=_T, top=_T, bottom=_T)


def _hoja(wb, nombre, df):
    ws = wb.create_sheet(nombre[:31])
    for j, c in enumerate(df.columns, 1):
        x = ws.cell(row=1, column=j, value=str(c))
        x.fill, x.font, x.border = HEAD, Font(name=F, bold=True, color="FFFFFF"), BD
        x.alignment = Alignment(horizontal="center", wrap_text=True)
    num = []
    for i, (_, fila) in enumerate(df.iterrows(), 2):
        for j, c in enumerate(df.columns, 1):
            v = fila[c]
            if pd.isna(v):
                v = None
            elif isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            elif hasattr(v, "item"):
                v = v.item()
            x = ws.cell(row=i, column=j, value=v)
            x.font, x.border = Font(name=F, size=10), BD
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                x.number_format = FMT
                if j not in num:
                    num.append(j)
            elif isinstance(v, datetime.datetime):
                x.number_format = "DD/MM/YYYY"
    tr = len(df) + 2
    ws.cell(row=tr, column=1, value="TOTAL").font = Font(name=F, bold=True)
    for j in num:
        L = get_column_letter(j)
        x = ws.cell(row=tr, column=j, value=f"=SUM({L}2:{L}{tr-1})")
        x.number_format, x.font = FMT, Font(name=F, bold=True)
    for j, c in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(j)].width = max(12, min(32, len(str(c)) + 5))
    ws.freeze_panes = "A2"


def generar(cfg, salida):
    wb = Workbook()
    ws = wb.active
    ws.title = "conciliacion"
    ws["B2"] = cfg.get("titulo", "Conciliación cuenta recaudación PedidosYa")
    ws["B2"].font = Font(name=F, bold=True, size=14)
    if cfg.get("subtitulo"):
        ws["B3"] = cfg["subtitulo"]
        ws["B3"].font = Font(name=F, italic=True, size=10)

    r = 5
    for j, h in enumerate(["Concepto", "Debe", "Haber", "Saldo", "Observaciones"]):
        x = ws.cell(row=r, column=2 + j, value=h)
        x.fill, x.font, x.border = HEAD, Font(name=F, bold=True, color="FFFFFF"), BD
        x.alignment = Alignment(horizontal="center")

    r += 1
    ws.cell(row=r, column=2, value="Saldo").font = Font(name=F, bold=True, size=11)
    x = ws.cell(row=r, column=5, value=cfg["saldo_inicial"])
    x.number_format, x.font, x.fill = FMT, Font(name=F, bold=True, size=11), SALDO
    prev = r

    for ln in cfg["lineas"]:
        r += 1
        if "subtotal" in ln:
            ws.cell(row=r, column=2, value=ln["subtotal"]).font = Font(name=F, bold=True, size=11)
            x = ws.cell(row=r, column=5, value=f"=E{prev}")
            x.number_format, x.font, x.fill = FMT, Font(name=F, bold=True, size=11), SALDO
            prev = r
            continue
        ws.cell(row=r, column=2, value=ln["concepto"]).font = Font(name=F, size=10)
        if ln.get("debe"):
            x = ws.cell(row=r, column=3, value=round(ln["debe"], 2))
            x.number_format, x.font = FMT, Font(name=F, size=10)
        if ln.get("haber"):
            x = ws.cell(row=r, column=4, value=round(ln["haber"], 2))
            x.number_format, x.font = FMT, Font(name=F, size=10)
        x = ws.cell(row=r, column=5, value=f"=E{prev}+C{r}-D{r}")
        x.number_format, x.font = FMT, Font(name=F, size=10)
        if ln.get("nota"):
            ws.cell(row=r, column=6, value=ln["nota"]).font = Font(name=F, size=9, italic=True)
        if ln.get("flag"):
            for k in range(2, 7):
                ws.cell(row=r, column=k).fill = FLAG
        prev = r

    r += 1
    ws.cell(row=r, column=2, value="SALDO FINAL").font = Font(name=F, bold=True, size=12)
    x = ws.cell(row=r, column=5, value=f"=E{prev}")
    x.number_format, x.font, x.fill = FMT, Font(name=F, bold=True, size=12), SALDO

    if cfg.get("abierto"):
        r += 2
        ws.cell(row=r, column=2, value="PENDIENTE DE DEFINICIÓN").font = Font(
            name=F, bold=True, size=10, color="C00000")
        for ln in cfg["abierto"]:
            r += 1
            ws.cell(row=r, column=2, value=ln["concepto"]).font = Font(name=F, size=10)
            for col, k in ((3, "debe"), (4, "haber")):
                if ln.get(k):
                    x = ws.cell(row=r, column=col, value=round(ln[k], 2))
                    x.number_format = FMT
            if ln.get("nota"):
                ws.cell(row=r, column=6, value=ln["nota"]).font = Font(name=F, size=9, italic=True)
            for k in range(2, 7):
                ws.cell(row=r, column=k).fill = FLAG

    if cfg.get("nota_final"):
        r += 2
        ws.cell(row=r, column=2, value=cfg["nota_final"]).font = Font(name=F, italic=True, size=9)

    for col, w in (("B", 68), ("C", 17), ("D", 17), ("E", 17), ("F", 44)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B6"

    # --- historico de diferencias por mes ---
    # Perdura entre cierres: cada mes agrega una columna. Es lo que permite no
    # re-subir el detalle crudo de los meses ya cerrados.
    if cfg.get("historico"):
        h = cfg["historico"]
        df = pd.read_csv(h) if isinstance(h, str) else pd.DataFrame(h)
        _hoja(wb, "Diferencias por mes", df)

    for nombre, ruta in (cfg.get("detalles") or {}).items():
        if not os.path.exists(ruta):
            print(f"  aviso: no existe {ruta}, se omite '{nombre}'")
            continue
        df = pd.read_csv(ruta) if ruta.lower().endswith(".csv") else pd.read_excel(ruta)
        _hoja(wb, nombre, df)

    wb.save(salida)
    return salida


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--json", required=True)
    ap.add_argument("--salida", default="Conciliacion.xlsx")
    a = ap.parse_args()
    ruta = generar(json.load(open(a.json, encoding="utf-8")), a.salida)
    print(f"Generado: {ruta}")
    print("Recalcular antes de leer valores:")
    print(f"  python /mnt/skills/public/xlsx/scripts/recalc.py {ruta}")


if __name__ == "__main__":
    main()
