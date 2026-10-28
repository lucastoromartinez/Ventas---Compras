"""Conciliación Rappi — submódulo EASA.

Parte del cruce de facturas contra liquidaciones (logica_rappi) y sigue desde
el cuadro de conceptos: cruza las ventas de Rappi contra Dean y HIO, cruza las
facturas y las acreditaciones contra el mayor de recaudación, y deja armado el
cuadro de pendientes del mes siguiente. Devuelve el mismo zip del cruce más el
libro 'Conciliacion.xlsx'.

Revisión 28/10/2026 (cierre de septiembre):
- El cuadro de pendientes puede traer todos los meses apilados: se trabaja
  sobre el último bloque y el resto se conserva como histórico.
- El mayor se recorta al último día del mes (venía con asientos de octubre).
- Las acreditaciones pendientes se leen sólo del bloque 'Acreditaciones de …
  Registradas en …' del último mes (antes tomaba líneas de meses viejos).
- Cruce de facturas por período: si el mayor tiene más importe que las facturas
  del período (otras ya registradas), o no hay combinación que explique la
  diferencia, la factura pasa al cruce por número en vez de quedar como
  'diferencia'.
- Formatos exigidos en el cuadro de pendientes: facturas
  'A-nnnn-nnnnnnnn Periodo dd-mm-aaaa al dd-mm-aaaa' y liquidaciones
  'Liquidacion dd/mm/aaaa a dd/mm/aaaa'. Otros formatos no se leen.
- Las facturas no asignadas que ya están en el pendientes, o que son de un
  período posterior al mes, no se vuelven a imputar.
- El bloque del mes se arma con una línea por concepto (sumado) en orden
  alfabético, como el cuadro manual; la 'Diferencia de centavos' va al pie del
  bloque 'Diferencia de Venta y extracto' del mes.
"""

import calendar
import datetime as dt
import io
import itertools
import re
import unicodedata
import zipfile
from collections import Counter

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill

from logica_rappi import (
    construir_stats_rappi,
    cuadro_conceptos_a_df,
    escribir_salidas_rappi,
    procesar_rappi,
)

MESES_ES = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
            "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]

TOLERANCIA_FACTURA = 0.01
TOLERANCIA_ACREDITACION = 5

MAX_FILAS_COMBINACION_COMPLETA = 15
MAX_TAM_COMBINACION_RECORTADA = 3


def _avisar(advertencias, mensaje):
    if advertencias is not None:
        advertencias.append(mensaje)


def _normalizar(texto):
    if texto is None:
        return ''
    sin_tildes = unicodedata.normalize('NFKD', str(texto))
    sin_tildes = ''.join(c for c in sin_tildes if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]', '', sin_tildes.lower())


def _buscar_col(df, *claves, obligatoria=True, nombre_humano=None):
    claves_norm = [_normalizar(c) for c in claves]
    for col in df.columns:
        col_norm = _normalizar(col)
        if all(k in col_norm for k in claves_norm):
            return col
    if obligatoria:
        raise ValueError(
            f"No se encontró la columna '{nombre_humano or ' '.join(claves)}'. "
            f"Columnas del archivo: {list(df.columns)}"
        )
    return None


def _renombrar_columnas(df, mapa):
    renombres = {}
    for canonico, claves in mapa.items():
        if canonico in df.columns:
            continue
        col = _buscar_col(df, *claves, obligatoria=True, nombre_humano=canonico)
        renombres[col] = canonico
    return df.rename(columns=renombres)


def _leer_bytes(archivo):
    if hasattr(archivo, 'getvalue'):
        return archivo.getvalue()
    try:
        archivo.seek(0)
    except (AttributeError, OSError):
        pass
    return archivo.read()


def _como_archivo(datos, nombre=None):
    buf = io.BytesIO(datos)
    if nombre:
        buf.name = nombre
    return buf


def _tiene_columnas(df, grupos_clave):
    return all(_buscar_col(df, *claves, obligatoria=False) is not None for claves in grupos_clave)


def _leer_reporte(archivo, etiqueta, grupos_clave=()):
    datos = _leer_bytes(archivo)
    df = pd.read_excel(_como_archivo(datos))

    if not grupos_clave or _tiene_columnas(df, grupos_clave):
        return df

    crudo = pd.read_excel(_como_archivo(datos), header=None, nrows=15)
    for i in range(len(crudo)):
        fila = [_normalizar(v) for v in crudo.iloc[i].tolist()]
        if all(any(all(_normalizar(k) in celda for k in claves) for celda in fila)
               for claves in grupos_clave):
            return pd.read_excel(_como_archivo(datos), header=i)

    esperadas = ', '.join(' '.join(claves) for claves in grupos_clave)
    raise ValueError(
        f"El archivo cargado en «{etiqueta}» no parece ser ese reporte: no tiene "
        f"las columnas esperadas ({esperadas}). Columnas encontradas: {list(df.columns)}. "
        "Revisá que no se haya mezclado con otro de los archivos."
    )


def _leer_pendientes(archivo, etiqueta):
    df = pd.read_excel(_como_archivo(_leer_bytes(archivo)), header=None)

    if df.shape[1] < 4:
        raise ValueError(
            f"El archivo cargado en «{etiqueta}» tiene {df.shape[1]} columna/s: el cuadro "
            "de pendientes necesita al menos 4 (el texto va en la segunda y el monto en la cuarta)."
        )

    texto = df.iloc[:, 1].apply(lambda x: "" if pd.isna(x) else _normalizar(x))
    if not texto.str.startswith("saldosegunmayorcontable").any():
        raise ValueError(
            f"El archivo cargado en «{etiqueta}» no parece el cuadro de pendientes: "
            "no tiene la línea 'Saldo según Mayor Contable'. "
            "Revisá que no se haya mezclado con otro de los archivos."
        )

    return df


def separar_historico_y_ultimo_bloque(df_pendientes, advertencias=None):
    """El archivo de conciliación acumula un bloque por mes (cada uno arranca
    con una fecha y 'Saldo según Mayor Contable al ...'). Se trabaja sólo sobre
    el último bloque; el resto queda como histórico y se vuelve a escribir tal
    cual adelante del bloque nuevo. Devuelve (historico, ultimo_bloque)."""
    texto = df_pendientes.iloc[:, 1].apply(lambda x: "" if pd.isna(x) else _normalizar(x))
    idx_saldo = [i for i, t in texto.items() if t.startswith("saldosegunmayorcontable")]
    inicio = idx_saldo[-1]
    # La fila de fecha que encabeza el bloque va con él.
    if inicio > 0 and isinstance(df_pendientes.iloc[inicio - 1, 1], (pd.Timestamp, dt.date, dt.datetime)):
        inicio -= 1
    elif inicio > 1 and df_pendientes.iloc[inicio - 1].isna().all() and \
            isinstance(df_pendientes.iloc[inicio - 2, 1], (pd.Timestamp, dt.date, dt.datetime)):
        inicio -= 2
    if len(idx_saldo) > 1:
        _avisar(advertencias,
                f"El cuadro de pendientes trae {len(idx_saldo)} meses apilados: se tomó el último "
                f"(fila {inicio + 1}) y los anteriores se conservan como histórico.")
    historico = df_pendientes.iloc[:inicio].reset_index(drop=True)
    ultimo = df_pendientes.iloc[inicio:].reset_index(drop=True)
    return historico, ultimo


def filtrar_mayor_hasta_fin_de_mes(recaudacion, anio, mes_numero, advertencias=None):
    """El mayor suele exportarse con asientos del mes siguiente (facturas que
    se registran en los primeros días de octubre, por ejemplo). Todo lo
    posterior al último día del mes que se concilia queda afuera: si no, el
    saldo final no es el del cierre y facturas todavía no registradas al cierre
    aparecen como cobradas."""
    df = recaudacion.copy()
    col_fecha = next((c for c in df.columns if _normalizar(c) == 'fecha'), None) or \
        _buscar_col(df, 'fecha', nombre_humano='Fecha')
    fechas = pd.to_datetime(df[col_fecha], errors='coerce', dayfirst=True)
    tope = pd.Timestamp(anio, mes_numero, calendar.monthrange(anio, mes_numero)[1])
    posteriores = fechas > tope
    if posteriores.any():
        _avisar(advertencias,
                f"Mayor: se excluyeron {int(posteriores.sum())} asiento/s posteriores al "
                f"{tope:%d/%m/%Y} (quedan para el mes siguiente).")
    return df[~posteriores].reset_index(drop=True)


def _normalizar_id(serie):
    def _conv(x):
        if pd.isna(x):
            return None
        s = str(x).strip()
        if re.match(r'^\d+\.0$', s):
            s = s.split('.')[0]
        return s
    return serie.apply(_conv)


def _hasta_fecha(serie, tope):
    fechas = pd.to_datetime(serie, errors='coerce')
    if fechas.empty:
        return pd.Series([], dtype=bool, index=serie.index)
    return fechas <= tope


def _buscar_subset(registros_df, objetivo_valor, col_valor="Importe", advertencias=None):
    registros = list(registros_df.itertuples())
    objetivo = round(abs(objetivo_valor), 2)

    tam_max = len(registros)
    if tam_max > MAX_FILAS_COMBINACION_COMPLETA:
        tam_max = MAX_TAM_COMBINACION_RECORTADA
        _avisar(advertencias,
                f"Hay {len(registros)} filas para explicar {objetivo}: se probaron "
                f"combinaciones de hasta {tam_max} filas para no colgar el proceso.")

    for tam in range(1, tam_max + 1):
        encontrados = [c for c in itertools.combinations(registros, tam)
                       if round(sum(getattr(r, col_valor) for r in c), 2) == objetivo]
        if encontrados:
            if len(encontrados) > 1:
                _avisar(advertencias,
                        f"{len(encontrados)} combinaciones de {tam} fila/s explican "
                        f"{objetivo}: se tomó la primera.")
            return [r.Index for r in encontrados[0]]
    return None


def importar_ventas_rappi(archivos_liq):
    dfs = []
    for nombre, datos in archivos_liq:
        try:
            df = pd.read_excel(_como_archivo(datos), sheet_name="Detalle", header=1)
        except ValueError:
            continue
        if "Tipo de transacción" not in df.columns:
            continue
        mask_orden = df["Tipo de transacción"].astype(str).str.strip().str.upper() == "ORDEN"
        df_filtrado = df[mask_orden].copy()
        df_filtrado["_archivo_origen"] = nombre
        dfs.append(df_filtrado)

    return pd.concat(dfs, ignore_index=True) if dfs else pd.DataFrame()


def depurar_venta_rappi(df):
    df = df.copy()

    meses = {'ene': '01', 'feb': '02', 'mar': '03', 'abr': '04',
             'may': '05', 'jun': '06', 'jul': '07', 'ago': '08',
             'sep': '09', 'oct': '10', 'nov': '11', 'dic': '12'}

    def parsear_fecha_hora(texto):
        if pd.isna(texto):
            return pd.NaT
        if isinstance(texto, pd.Timestamp):
            return texto

        texto = str(texto).strip().lower()
        primer_digito = re.search(r'\d', texto)
        if not primer_digito:
            return pd.NaT
        texto_recortado = texto[primer_digito.start():]

        m = re.match(
            r'(\d{1,2})\s+([a-zñáéíóú]+)\.\s+(\d{4}),\s+(\d{1,2}):(\d{2}):(\d{2})\s*(a|p)\.\s*m\.',
            texto_recortado
        )
        if not m:
            return pd.to_datetime(texto, errors="coerce")

        dia, mes_abr, anio, hora, minuto, segundo, periodo = m.groups()
        mes = meses[mes_abr[:3]]
        hora = int(hora)
        if periodo == 'p' and hora != 12:
            hora += 12
        if periodo == 'a' and hora == 12:
            hora = 0
        return pd.Timestamp(f"{anio}-{mes}-{dia.zfill(2)} {hora:02d}:{minuto}:{segundo}")

    fecha_hora = df["Fecha de creación orden"].apply(parsear_fecha_hora)

    idx_col = df.columns.get_loc("Fecha de creación orden")
    df = df.drop(columns=["Fecha de creación orden"])
    df.insert(idx_col, "Hora", fecha_hora.dt.strftime("%H:%M:%S"))
    df.insert(idx_col, "Fecha", fecha_hora.dt.normalize())

    def a_objeto_id(serie):
        return serie.apply(lambda x: str(int(x)) if pd.notna(x) else None).astype(object)

    for col in ["ID de la órden", "ID del paidlot", "ID de la tienda"]:
        if col in df.columns:
            df[col] = a_objeto_id(df[col])

    for col in df.select_dtypes(include="float64").columns:
        df[col] = df[col].round(2)

    idx_venta_bruta = df.columns.get_loc("Venta Bruta")
    df.insert(
        idx_venta_bruta + 1,
        "Venta Neta",
        (df["Venta Bruta"] + df["Descuento de Producto asumido por el aliado"]).round(2)
    )

    return df


def depurar_dean_rappi(df, advertencias=None):
    df = _renombrar_columnas(df.copy(), {
        'Fecha':    ('fecha',),
        'Online':   ('online',),
        'Efectivo': ('efectivo',),
    })

    fecha_hora = pd.to_datetime(df["Fecha"], dayfirst=True, errors="coerce")

    idx_col = df.columns.get_loc("Fecha")
    df = df.drop(columns=["Fecha"])
    df.insert(idx_col, "Hora", fecha_hora.dt.strftime("%H:%M:%S"))
    df.insert(idx_col, "Fecha", fecha_hora.dt.normalize())

    filas_antes = len(df)
    df = df[df["Fecha"].notna()].reset_index(drop=True)
    if filas_antes != len(df):
        _avisar(advertencias, f"Dean: se descartaron {filas_antes - len(df)} filas sin fecha.")

    idx_efectivo = df.columns.get_loc("Efectivo")
    df.insert(idx_efectivo + 1, "Importe", df["Online"].fillna(0) + df["Efectivo"].fillna(0))

    for col in df.select_dtypes(include="float64").columns:
        df[col] = df[col].round(2)

    return df


def cruzar_hio_documento_metodo(df_hio_documento, df_hio_metodo):
    df_hio_metodo = _renombrar_columnas(df_hio_metodo.copy(), {
        'Serie / Número': ('serie',),
        'Medio Pago':     ('medio', 'pago'),
    })
    df_hio_documento = _renombrar_columnas(df_hio_documento.copy(), {
        'Serie / Número': ('serie',),
    })

    df_hio_metodo = df_hio_metodo[
        df_hio_metodo["Medio Pago"].astype(str).str.upper().str.contains("RAPPI", na=False)
    ].copy()

    df_hio_documento["Serie / Número"] = df_hio_documento["Serie / Número"].astype(str).str.strip()
    df_hio_metodo["Serie / Número"] = df_hio_metodo["Serie / Número"].astype(str).str.strip()

    df_hio_documento["_matched"] = False
    df_hio_metodo["_matched"] = False

    for i, row_doc in df_hio_documento.iterrows():
        candidatos = df_hio_metodo[
            (df_hio_metodo["Serie / Número"] == row_doc["Serie / Número"]) &
            (~df_hio_metodo["_matched"])
        ]
        if not candidatos.empty:
            j = candidatos.index[0]
            df_hio_documento.at[i, "_matched"] = True
            df_hio_metodo.at[j, "_matched"] = True

    match_documento = df_hio_documento[df_hio_documento["_matched"]].drop(columns="_matched")
    falta_documento = df_hio_documento[~df_hio_documento["_matched"]].drop(columns="_matched")
    falta_metodo = df_hio_metodo[~df_hio_metodo["_matched"]].drop(columns="_matched")

    return match_documento, falta_documento, falta_metodo


def cruzar_dean(df_venta_rappi_dep, df_dean_dep):
    df_rappi_dean = df_venta_rappi_dep[
        df_venta_rappi_dep["Nombre de la tienda"].astype(str).str.contains("Dean", case=False, na=False)
    ].copy()
    df_dean = _renombrar_columnas(df_dean_dep.copy(), {
        'Nº Transacción': ('transacc',),
    })

    df_rappi_dean["ID de la órden"] = _normalizar_id(df_rappi_dean["ID de la órden"])
    df_dean["Nº Transacción"] = _normalizar_id(df_dean["Nº Transacción"])

    df_rappi_dean["_matched"] = False
    df_dean["_matched"] = False
    df_rappi_dean["Importe Dean"] = np.nan

    for i, row_dean in df_dean.iterrows():
        candidatos = df_rappi_dean[
            (df_rappi_dean["ID de la órden"] == row_dean["Nº Transacción"]) &
            (~df_rappi_dean["_matched"])
        ]
        if not candidatos.empty:
            j = candidatos.index[0]
            df_dean.at[i, "_matched"] = True
            df_rappi_dean.at[j, "_matched"] = True
            df_rappi_dean.at[j, "Importe Dean"] = row_dean["Importe"]

    df_rappi_dean["Diferencia"] = (
        df_rappi_dean["Importe Dean"] - df_rappi_dean["Venta Neta"]
    ).round(2)

    match_dean1 = df_dean[df_dean["_matched"]].drop(columns="_matched")
    match_rappi_dean1 = df_rappi_dean[df_rappi_dean["_matched"]].drop(columns="_matched")
    falta_rappi_dean1 = df_dean[~df_dean["_matched"]].drop(columns="_matched")
    falta_dean1 = df_rappi_dean[~df_rappi_dean["_matched"]].drop(
        columns=["_matched", "Importe Dean", "Diferencia"])

    return match_dean1, match_rappi_dean1, falta_rappi_dean1, falta_dean1


def cruzar_hio(df_hio, df_venta_rappi_dep):
    df_rappi_hio = df_venta_rappi_dep[
        ~df_venta_rappi_dep["Nombre de la tienda"].astype(str).str.contains(
            "dean|atalaya", case=False, na=False, regex=True
        )
    ].copy()
    df_hio = _renombrar_columnas(df_hio.copy(), {
        'Localizador': ('localizador',),
        'Venta':       ('venta',),
    })

    df_rappi_hio["ID de la órden"] = _normalizar_id(df_rappi_hio["ID de la órden"])
    df_hio["Localizador"] = _normalizar_id(df_hio["Localizador"])

    df_rappi_hio["_matched"] = False
    df_hio["_matched"] = False
    df_rappi_hio["Importe Hio"] = np.nan

    for i, row_hio in df_hio.iterrows():
        candidatos = df_rappi_hio[
            (df_rappi_hio["ID de la órden"] == row_hio["Localizador"]) &
            (~df_rappi_hio["_matched"])
        ]
        if not candidatos.empty:
            j = candidatos.index[0]
            df_hio.at[i, "_matched"] = True
            df_rappi_hio.at[j, "_matched"] = True
            df_rappi_hio.at[j, "Importe Hio"] = row_hio["Venta"]

    df_rappi_hio["Diferencia"] = (
        df_rappi_hio["Importe Hio"] - df_rappi_hio["Venta Neta"]
    ).round(2)

    match_hio = df_hio[df_hio["_matched"]].drop(columns="_matched")
    match_hio_rappi1 = df_rappi_hio[df_rappi_hio["_matched"]].drop(columns="_matched")
    falta_rappi_hio1 = df_hio[~df_hio["_matched"]].drop(columns="_matched")
    falta_hio1 = df_rappi_hio[~df_rappi_hio["_matched"]].drop(
        columns=["_matched", "Importe Hio", "Diferencia"])

    return match_hio, match_hio_rappi1, falta_rappi_hio1, falta_hio1


def cruzar_hio_fecha_venta(falta_rappi_hio1, falta_hio1):
    def _hora_a_timedelta(serie):
        def _conv(x):
            if pd.isna(x):
                return pd.NaT
            if isinstance(x, dt.time):
                return pd.Timedelta(hours=x.hour, minutes=x.minute, seconds=x.second)
            return pd.to_timedelta(str(x), errors="coerce")
        return serie.apply(_conv)

    df_hio = falta_rappi_hio1.copy()
    df_rappi = falta_hio1.copy()

    df_hio["Fecha"] = pd.to_datetime(df_hio["Fecha"], errors="coerce").dt.normalize()
    df_rappi["Fecha"] = pd.to_datetime(df_rappi["Fecha"], errors="coerce").dt.normalize()

    df_hio["Venta"] = df_hio["Venta"].round(2)
    df_rappi["Venta Neta"] = df_rappi["Venta Neta"].round(2)

    df_hio["_hora_td"] = _hora_a_timedelta(df_hio["Hora"])
    df_rappi["_hora_td"] = _hora_a_timedelta(df_rappi["Hora"])

    df_hio["_matched"] = False
    df_rappi["_matched"] = False
    df_rappi["Importe Hio"] = np.nan

    for i, row_hio in df_hio.iterrows():
        candidatos = df_rappi[
            (df_rappi["Fecha"] == row_hio["Fecha"]) &
            (df_rappi["Venta Neta"] == row_hio["Venta"]) &
            (~df_rappi["_matched"])
        ]
        if candidatos.empty:
            continue
        if len(candidatos) == 1:
            j = candidatos.index[0]
        else:
            diffs = (candidatos["_hora_td"] - row_hio["_hora_td"]).abs()
            j = diffs.idxmin()

        df_hio.at[i, "_matched"] = True
        df_rappi.at[j, "_matched"] = True
        df_rappi.at[j, "Importe Hio"] = row_hio["Venta"]

    df_rappi["Diferencia"] = (df_rappi["Importe Hio"] - df_rappi["Venta Neta"]).round(2)

    match_hio2 = df_hio[df_hio["_matched"]].drop(columns=["_matched", "_hora_td"])
    match_hio_rappi2 = df_rappi[df_rappi["_matched"]].drop(columns=["_matched", "_hora_td"])
    falta_rappi_hio2 = df_hio[~df_hio["_matched"]].drop(columns=["_matched", "_hora_td"])
    falta_hio2 = df_rappi[~df_rappi["_matched"]].drop(
        columns=["_matched", "_hora_td", "Importe Hio", "Diferencia"])

    return match_hio2, match_hio_rappi2, falta_rappi_hio2, falta_hio2


def calcular_total_dean(df_venta_rappi_dep, df_dean_dep):
    rappi_dean = df_venta_rappi_dep[
        df_venta_rappi_dep["Nombre de la tienda"].astype(str).str.contains("dean", case=False, na=False)
    ]
    if rappi_dean.empty and df_dean_dep.empty:
        return None
    return round(rappi_dean["Venta Neta"].sum() - df_dean_dep["Importe"].sum(), 2)


def calcular_total_hio(match_hio_rappi1, falta_rappi_hio2, falta_hio2):
    if match_hio_rappi1.empty and falta_rappi_hio2.empty and falta_hio2.empty:
        return None
    diferencia_localizador = match_hio_rappi1["Diferencia"].sum()
    total_falta_hio = falta_rappi_hio2["Venta"].sum()
    total_falta_rappi = falta_hio2["Venta Neta"].sum()
    return round(-diferencia_localizador + total_falta_hio - total_falta_rappi, 2)


def calcular_venta_dean(df_dean_dep):
    return round(df_dean_dep["Importe"].sum(), 2)


def calcular_venta_hio(df_hio):
    return round(df_hio["Venta"].sum(), 2)


PATRON_FACTURA_NUM = re.compile(r'^A-\d{4}-\d{8}$')
PATRON_FACTURA_CTA_DIA = re.compile(r'^[A-Za-z]+\d+$')
PATRON_LIQUIDACION = re.compile(r'^Liquidacion (\d{2})/(\d{2})/(\d{4}) a (\d{2})/(\d{2})/(\d{4})$')
PATRON_FECHA_PAGO = re.compile(r'^Fecha de Pago (\d{2})/(\d{2})/(\d{4})$')
# '<Nro Factura> Periodo dd-mm-aaaa al dd-mm-aaaa' (como se escribe en el pendientes)
PATRON_FACTURA_CON_PERIODO = re.compile(
    r'^(A-\d{4}-\d{8}|[A-Za-z]+\d+)\s+Periodo\s+(?:del\s+)?(\d{2}-\d{2}-\d{4})\s+al\s+(\d{2}-\d{2}-\d{4})$'
)


def _texto_concepto(row):
    return str(row["Concepto"]).strip() if pd.notna(row["Concepto"]) else ""


def extraer_facturas_cuadro_conceptos(df_cuadro_conceptos):
    periodo_actual, fecha_pago_actual = None, None
    filas = []

    for idx, row in df_cuadro_conceptos.iterrows():
        texto = _texto_concepto(row)
        monto = row["Monto"]

        if texto == "Facturas no asignadas":
            periodo_actual, fecha_pago_actual = None, None
            continue
        if texto == "":
            continue

        m_liq = PATRON_LIQUIDACION.match(texto)
        if m_liq:
            d1, m1, a1, d2, m2, a2 = m_liq.groups()
            periodo_actual = f"Periodo {d1}-{m1}-{a1} al {d2}-{m2}-{a2}"
            continue

        m_pago = PATRON_FECHA_PAGO.match(texto)
        if m_pago:
            dp, mp, ap = m_pago.groups()
            fecha_pago_actual = pd.Timestamp(int(ap), int(mp), int(dp))
            continue

        # Una factura no asignada puede venir con su período escrito al lado
        # ('A-0002-02046013 Periodo 01-10-2026 al 04-10-2026'): se usa ese.
        m_fac = PATRON_FACTURA_CON_PERIODO.match(texto)
        if m_fac:
            nro, periodo_propio = m_fac.group(1), f"Periodo {m_fac.group(2)} al {m_fac.group(3)}"
        else:
            nro, periodo_propio = texto, None

        if PATRON_FACTURA_NUM.match(nro) or PATRON_FACTURA_CTA_DIA.match(nro):
            formato = "numerico" if PATRON_FACTURA_NUM.match(nro) else "cuentas_al_dia"
            filas.append({
                "Orden": idx,
                "Nro Factura": nro,
                "Importe": round(float(monto), 2),
                "Formato": formato,
                "Periodo": periodo_actual or periodo_propio,
                "Fecha de Pago": fecha_pago_actual,
                "Asignada": periodo_actual is not None,
            })

    return pd.DataFrame(
        filas, columns=["Orden", "Nro Factura", "Importe", "Formato", "Periodo", "Fecha de Pago", "Asignada"]
    ).reset_index(drop=True)


def depurar_facturas_no_asignadas(facturas_periodo_actual, facturas_pendientes, anio, mes_numero,
                                  advertencias=None):
    """Las facturas que no matchearon con ninguna liquidación del mes pueden ser
    de dos clases que no corresponden a este cierre:
    - las que ya figuran en el cuadro de pendientes del mes anterior (se cruzan
      por ese lado; si quedaran acá también, se contarían dos veces);
    - las de períodos posteriores al mes (se subieron facturas de octubre al
      conciliar septiembre): van al mes que viene.
    Las demás (sin liquidación y sin período posterior) se dejan, con aviso."""
    df = facturas_periodo_actual.copy()
    if df.empty:
        return df

    ya_pendientes = set(facturas_pendientes["Nro Factura"]) if facturas_pendientes is not None else set()
    tope = pd.Timestamp(anio, mes_numero, calendar.monthrange(anio, mes_numero)[1])

    def _inicio_periodo(p):
        m = re.match(r'^Periodo (\d{2})-(\d{2})-(\d{4})', str(p))
        return pd.Timestamp(int(m.group(3)), int(m.group(2)), int(m.group(1))) if m else pd.NaT

    no_asignada = ~df["Asignada"].astype(bool)
    repetida = no_asignada & df["Nro Factura"].isin(ya_pendientes)
    posterior = no_asignada & (df["Periodo"].apply(_inicio_periodo) > tope)
    sin_dato = no_asignada & ~repetida & ~posterior

    if repetida.any():
        _avisar(advertencias,
                f"Facturas no asignadas que ya estaban en el cuadro de pendientes (se cruzan por ahí): "
                f"{', '.join(df.loc[repetida, 'Nro Factura'])}.")
    if posterior.any():
        _avisar(advertencias,
                f"Facturas de períodos posteriores al {tope:%d/%m/%Y}, se dejan para el mes siguiente: "
                f"{', '.join(df.loc[posterior, 'Nro Factura'])}.")
    if sin_dato.any():
        _avisar(advertencias,
                f"Facturas sin liquidación del mes (se cruzan igual contra el mayor): "
                f"{', '.join(df.loc[sin_dato, 'Nro Factura'])}.")

    return df[~repetida & ~posterior].reset_index(drop=True)


def extraer_conceptos_cuadro_conceptos(df_cuadro_conceptos):
    tienda_actual, periodo_actual, fecha_pago_actual = None, None, None
    filas = []

    for idx, row in df_cuadro_conceptos.iterrows():
        texto = _texto_concepto(row)
        monto = row["Monto"]

        if texto == "Facturas no asignadas":
            tienda_actual, periodo_actual, fecha_pago_actual = None, None, None
            continue
        if texto == "":
            continue

        m_liq = PATRON_LIQUIDACION.match(texto)
        if m_liq:
            d1, m1, a1, d2, m2, a2 = m_liq.groups()
            periodo_actual = f"Periodo {d1}-{m1}-{a1} al {d2}-{m2}-{a2}"
            continue

        m_pago = PATRON_FECHA_PAGO.match(texto)
        if m_pago:
            dp, mp, ap = m_pago.groups()
            fecha_pago_actual = pd.Timestamp(int(ap), int(mp), int(dp))
            continue

        if pd.isna(monto):
            tienda_actual = texto
            continue

        if texto == "Valor total a transferir":
            continue
        if PATRON_FACTURA_NUM.match(texto) or PATRON_FACTURA_CTA_DIA.match(texto) \
                or PATRON_FACTURA_CON_PERIODO.match(texto):
            continue

        filas.append({
            "Orden": idx,
            "Concepto": texto,
            "Monto": round(float(monto), 2),
            "Tienda": tienda_actual,
            "Periodo": periodo_actual,
            "Fecha de Pago": fecha_pago_actual,
        })

    return pd.DataFrame(
        filas, columns=["Orden", "Concepto", "Monto", "Tienda", "Periodo", "Fecha de Pago"]
    ).reset_index(drop=True)


def extraer_liquidaciones_cuadro_conceptos(df_cuadro_conceptos):
    tienda_actual, concepto_actual, fecha_pago_actual = None, None, None
    filas = []

    for idx, row in df_cuadro_conceptos.iterrows():
        texto = _texto_concepto(row)
        monto = row["Monto"]

        if texto == "Facturas no asignadas":
            break
        if texto == "":
            continue

        if PATRON_LIQUIDACION.match(texto):
            concepto_actual = texto
            continue

        m_pago = PATRON_FECHA_PAGO.match(texto)
        if m_pago:
            dp, mp, ap = m_pago.groups()
            fecha_pago_actual = pd.Timestamp(int(ap), int(mp), int(dp))
            continue

        if texto == "Valor total a transferir":
            filas.append({
                "Orden": idx,
                "Tienda": tienda_actual,
                "Concepto": concepto_actual,
                "Fecha de Pago": fecha_pago_actual,
                "Importe": round(float(monto), 2),
            })
            continue

        if pd.isna(monto):
            tienda_actual = texto

    return pd.DataFrame(
        filas, columns=["Orden", "Tienda", "Concepto", "Fecha de Pago", "Importe"]
    ).reset_index(drop=True)


def detectar_mes_cuadro_conceptos(df_cuadro_conceptos, advertencias=None):
    meses = []
    for _, row in df_cuadro_conceptos.iterrows():
        m = PATRON_LIQUIDACION.match(_texto_concepto(row))
        if m:
            meses.append((int(m.group(3)), int(m.group(2))))

    if not meses:
        return None, None

    (anio, mes), cantidad = Counter(meses).most_common(1)[0]
    if len(set(meses)) > 1:
        _avisar(advertencias,
                f"El cuadro tiene liquidaciones de {len(set(meses))} meses distintos: "
                f"se tomó {mes:02d}/{anio} ({cantidad} de {len(meses)} liquidaciones).")
    return anio, mes


def _resolver_mes(anio, mes_numero, cuadro_conceptos, advertencias=None):
    if anio is not None and mes_numero is not None:
        return anio, mes_numero
    if cuadro_conceptos is None:
        raise ValueError("Pasá anio y mes_numero, o cuadro_conceptos para detectarlos.")
    anio, mes_numero = detectar_mes_cuadro_conceptos(cuadro_conceptos, advertencias)
    if anio is None:
        raise ValueError("No se pudo detectar el mes: el cuadro no tiene liquidaciones.")
    return anio, mes_numero


def _preparar_recaudacion(recaudacion):
    df_rec = _renombrar_columnas(recaudacion.copy(), {
        'Comentario':  ('comentario',),
        'Haber':       ('haber',),
        'Su factura':  ('sufactura',),
    })
    df_rec["Su factura"] = df_rec["Su factura"].astype(str).str.strip()
    df_rec.loc[df_rec["Su factura"].isin(["nan", "None", ""]), "Su factura"] = None
    return df_rec


def _homologar_factura(nro):
    m = re.match(r'^A-(\d{4}-\d{8})$', str(nro).strip())
    return m.group(1) if m else None


def cruzar_facturas_periodo_actual(recaudacion, facturas_periodo_actual, advertencias=None):
    def _digitos(texto):
        return re.sub(r'\D', '', str(texto))

    sin_periodo = facturas_periodo_actual[facturas_periodo_actual["Periodo"].isna()].copy()
    con_periodo = facturas_periodo_actual[facturas_periodo_actual["Periodo"].notna()].copy()

    df_rec = _preparar_recaudacion(recaudacion)

    es_periodo_txt = df_rec["Comentario"].astype(str).str.strip().str.lower().str.startswith("periodo")
    resumen = df_rec[es_periodo_txt].groupby("Comentario", as_index=False)["Haber"].sum()
    resumen["Haber"] = resumen["Haber"].round(2)
    resumen["_digitos"] = resumen["Comentario"].apply(_digitos)

    def _buscar_haber_periodo(digitos_periodo):
        c = resumen[resumen["_digitos"].apply(lambda d: bool(d) and digitos_periodo.startswith(d))]
        return c["Haber"].sum() if not c.empty else None

    filas_match, filas_diferencia, candidatas = [], [], []

    for periodo, grupo in con_periodo.groupby("Periodo"):
        importe_total = round(grupo["Importe"].sum(), 2)
        haber = _buscar_haber_periodo(_digitos(periodo))
        if haber is None:
            candidatas.append(grupo)
            continue
        diff = round(haber - importe_total, 2)
        if abs(diff) <= TOLERANCIA_FACTURA:
            filas_match.append({"Periodo": periodo, "Nro Factura": None,
                                "Importe": importe_total, "Haber": haber})
            continue
        # Si el mayor tiene MÁS que las facturas del grupo, el excedente es de
        # otras facturas del mismo período ya registradas: el total por período
        # no sirve y se pasa al cruce por número. Si tiene MENOS, se busca qué
        # facturas faltan; si ninguna combinación lo explica, también se pasa
        # al cruce por número en lugar de darlo por diferencia.
        idx_faltante = _buscar_subset(grupo, diff, advertencias=advertencias) if diff < 0 else None
        if idx_faltante is None:
            candidatas.append(grupo)
            continue
        candidatas.append(grupo.loc[idx_faltante])
        resto = grupo.drop(index=idx_faltante)
        if not resto.empty:
            filas_match.append({"Periodo": periodo, "Nro Factura": None,
                                "Importe": round(resto["Importe"].sum(), 2), "Haber": haber})

    candidatas.append(sin_periodo)
    df_candidatas = pd.concat(candidatas, ignore_index=True) if candidatas else \
        pd.DataFrame(columns=facturas_periodo_actual.columns)

    haber_por_factura = df_rec.dropna(subset=["Su factura"]).groupby("Su factura")["Haber"].sum()

    pre_falta = []
    for _, fila in df_candidatas.iterrows():
        key = _homologar_factura(fila["Nro Factura"])
        if key is None or key not in haber_por_factura.index:
            pre_falta.append({"Orden": fila["Orden"], "Periodo": fila["Periodo"],
                              "Nro Factura": fila["Nro Factura"], "Importe": fila["Importe"],
                              "Fecha de Pago": fila["Fecha de Pago"]})
            continue
        haber = round(haber_por_factura[key], 2)
        diff = round(haber - fila["Importe"], 2)
        if abs(diff) <= TOLERANCIA_FACTURA:
            filas_match.append({"Periodo": fila["Periodo"], "Nro Factura": fila["Nro Factura"],
                                "Importe": fila["Importe"], "Haber": haber})
        else:
            filas_diferencia.append({"Periodo": fila["Periodo"], "Nro Factura": fila["Nro Factura"],
                                     "Diferencia": diff})

    agregados = (df_rec[es_periodo_txt & df_rec["Su factura"].isna()]
                 .groupby("Comentario", as_index=False)["Haber"].sum())
    agregados["Haber"] = agregados["Haber"].round(2)

    pool = pd.DataFrame(pre_falta, columns=["Orden", "Periodo", "Nro Factura", "Importe", "Fecha de Pago"])
    usados = set()

    if not pool.empty:
        for _, fila_agg in agregados.iterrows():
            disponibles = pool[~pool.index.isin(usados)]
            if disponibles.empty:
                break
            idx_match = _buscar_subset(disponibles, fila_agg["Haber"],
                                       col_valor="Importe", advertencias=advertencias)
            if idx_match is None:
                continue
            usados.update(idx_match)
            for _, f in pool.loc[idx_match].iterrows():
                filas_match.append({"Periodo": f["Periodo"], "Nro Factura": f["Nro Factura"],
                                    "Importe": f["Importe"], "Haber": f["Importe"]})

    falta = pool[~pool.index.isin(usados)].reset_index(drop=True)
    match = pd.DataFrame(filas_match, columns=["Periodo", "Nro Factura", "Importe", "Haber"])
    diferencia = pd.DataFrame(filas_diferencia, columns=["Periodo", "Nro Factura", "Diferencia"])

    return match, falta, diferencia


def extraer_facturas_pendientes(df_pendientes):
    patron_completo = re.compile(
        r'^(A-\d{4}-\d{8}|[A-Za-z]+\d+)\s+Periodo\s+(?:del\s+)?(\d{2}-\d{2}-\d{4})\s+al\s+(\d{2}-\d{2}-\d{4})$'
    )

    texto = df_pendientes.iloc[:, 1].apply(lambda x: "" if pd.isna(x) else str(x).strip())
    monto = pd.to_numeric(df_pendientes.iloc[:, 3], errors="coerce")

    filas = []
    for idx, t in texto.items():
        m = patron_completo.match(t)
        if not m or pd.isna(monto[idx]):
            continue
        nro_factura, periodo = m.group(1), f"Periodo {m.group(2)} al {m.group(3)}"
        formato = "numerico" if PATRON_FACTURA_NUM.match(nro_factura) else (
            "cuentas_al_dia" if PATRON_FACTURA_CTA_DIA.match(nro_factura) else None
        )
        filas.append({
            "Nro Factura": nro_factura,
            "Periodo": periodo,
            "Importe": round(abs(monto[idx]), 2),
            "Formato": formato,
        })

    return pd.DataFrame(filas, columns=["Nro Factura", "Periodo", "Importe", "Formato"]).reset_index(drop=True)


def cruzar_facturas_pendientes(recaudacion, facturas_pendientes, advertencias=None):
    patron_fecha = re.compile(r'(\d{1,2})[/-](\d{1,2})(?:[/-]\d{2,4})?')

    def _clave_periodo(texto):
        fechas = patron_fecha.findall(str(texto))
        if len(fechas) < 2:
            return None
        (d1, m1), (d2, m2) = fechas[0], fechas[1]
        return f"{int(d1):02d}{int(m1):02d}_{int(d2):02d}{int(m2):02d}"

    df_rec = _preparar_recaudacion(recaudacion)
    df_rec["_clave_periodo"] = df_rec["Comentario"].apply(_clave_periodo)

    resumen_recaudacion = (
        df_rec.dropna(subset=["_clave_periodo"])
        .groupby("_clave_periodo", as_index=False)["Haber"].sum()
    )
    resumen_recaudacion["Haber"] = resumen_recaudacion["Haber"].round(2)
    haber_por_clave = resumen_recaudacion.set_index("_clave_periodo")["Haber"]

    filas_match, filas_diferencia, candidatas_etapa2 = [], [], []

    for periodo, grupo in facturas_pendientes.groupby("Periodo"):
        importe_total = round(grupo["Importe"].sum(), 2)
        clave = _clave_periodo(periodo)
        haber = haber_por_clave.get(clave) if clave else None

        if haber is None:
            candidatas_etapa2.append(grupo)
            continue

        diff = round(haber - importe_total, 2)
        if abs(diff) <= TOLERANCIA_FACTURA:
            filas_match.append({"Periodo": periodo, "Nro Factura": None,
                                "Importe": importe_total, "Haber": haber})
            continue

        # Si el mayor tiene MÁS que las facturas del grupo, el excedente es de
        # otras facturas del mismo período ya registradas: el total por período
        # no sirve y se pasa al cruce por número. Si tiene MENOS, se busca qué
        # facturas faltan; si ninguna combinación lo explica, también se pasa
        # al cruce por número en lugar de darlo por diferencia.
        idx_faltante = _buscar_subset(grupo, diff, advertencias=advertencias) if diff < 0 else None
        if idx_faltante is None:
            candidatas_etapa2.append(grupo)
            continue

        candidatas_etapa2.append(grupo.loc[idx_faltante])
        resto = grupo.drop(index=idx_faltante)
        if not resto.empty:
            filas_match.append({"Periodo": periodo, "Nro Factura": None,
                                "Importe": round(resto["Importe"].sum(), 2), "Haber": haber})

    df_candidatas = pd.concat(candidatas_etapa2, ignore_index=True) if candidatas_etapa2 else \
        pd.DataFrame(columns=["Nro Factura", "Periodo", "Importe", "Formato"])

    haber_por_factura = df_rec.dropna(subset=["Su factura"]).groupby("Su factura")["Haber"].sum()

    filas_falta_pre_etapa3 = []
    for _, fila in df_candidatas.iterrows():
        key = _homologar_factura(fila["Nro Factura"])
        if key is None or key not in haber_por_factura.index:
            filas_falta_pre_etapa3.append({"Periodo": fila["Periodo"],
                                           "Nro Factura": fila["Nro Factura"],
                                           "Importe": fila["Importe"]})
            continue

        haber = round(haber_por_factura[key], 2)
        diff = round(haber - fila["Importe"], 2)
        if abs(diff) <= TOLERANCIA_FACTURA:
            filas_match.append({"Periodo": fila["Periodo"], "Nro Factura": fila["Nro Factura"],
                                "Importe": fila["Importe"], "Haber": haber})
        else:
            filas_diferencia.append({"Periodo": fila["Periodo"], "Nro Factura": fila["Nro Factura"],
                                     "Diferencia": diff})

    agregados = (
        df_rec[df_rec["_clave_periodo"].notna() & df_rec["Su factura"].isna()]
        .groupby("_clave_periodo", as_index=False)["Haber"].sum()
    )
    agregados["Haber"] = agregados["Haber"].round(2)

    pool_falta = pd.DataFrame(filas_falta_pre_etapa3, columns=["Periodo", "Nro Factura", "Importe"])
    usados = set()

    if not pool_falta.empty:
        for _, fila_agg in agregados.iterrows():
            disponibles = pool_falta[~pool_falta.index.isin(usados)]
            if disponibles.empty:
                break
            idx_match = _buscar_subset(disponibles, fila_agg["Haber"],
                                       col_valor="Importe", advertencias=advertencias)
            if idx_match is None:
                continue
            usados.update(idx_match)
            for _, f in pool_falta.loc[idx_match].iterrows():
                filas_match.append({"Periodo": f["Periodo"], "Nro Factura": f["Nro Factura"],
                                    "Importe": f["Importe"], "Haber": f["Importe"]})

    falta = pool_falta[~pool_falta.index.isin(usados)].reset_index(drop=True)
    match = pd.DataFrame(filas_match, columns=["Periodo", "Nro Factura", "Importe", "Haber"])
    diferencia = pd.DataFrame(filas_diferencia, columns=["Periodo", "Nro Factura", "Diferencia"])

    return match, falta, diferencia


def extraer_acreditaciones_pendientes(pendientes_mes_anterior):
    """Sólo las liquidaciones del bloque 'Acreditaciones de <mes> Registradas
    en <mes siguiente>' (no cualquier 'Liquidacion ...' que ande por el cuadro),
    en el formato 'Liquidacion dd/mm/aaaa a dd/mm/aaaa'."""
    patron_header = re.compile(r'^Acreditaciones de .+ Registradas en .+$', re.IGNORECASE)

    texto = pendientes_mes_anterior.iloc[:, 1].apply(lambda x: "" if pd.isna(x) else str(x).strip())
    monto = pd.to_numeric(pendientes_mes_anterior.iloc[:, 3], errors="coerce")

    filas, dentro = [], False
    for idx, t in texto.items():
        if patron_header.match(t) and not t.lower().startswith("gastos"):
            dentro = True
            continue
        if not dentro:
            continue
        if PATRON_LIQUIDACION.match(t) and pd.notna(monto[idx]):
            filas.append({"Concepto": t, "Importe": round(abs(monto[idx]), 2)})
        else:
            dentro = False
    return pd.DataFrame(filas, columns=["Concepto", "Importe"]).reset_index(drop=True)


def separar_acreditaciones_negativas(liquidaciones_mes_actual):
    negativas = liquidaciones_mes_actual[liquidaciones_mes_actual["Importe"] < 0].reset_index(drop=True)
    positivas = liquidaciones_mes_actual[liquidaciones_mes_actual["Importe"] >= 0].reset_index(drop=True)
    return positivas, negativas


def cruzar_acreditaciones(recaudacion, liquidaciones_mes_actual, acred_mes_anterior,
                          anio, mes_numero, advertencias=None):
    patron_acred = re.compile(rf'acreditaci[oó]n\s+rappi\s+{mes_numero:02d}[-/]{anio}', re.IGNORECASE)
    df_rec = _renombrar_columnas(recaudacion.copy(), {
        'Comentario': ('comentario',),
        'Haber':      ('haber',),
    })
    comentario = df_rec["Comentario"].astype(str)
    total_mayor = round(df_rec.loc[comentario.str.contains(patron_acred, na=False), "Haber"].sum(), 2)

    ultimo_dia_mes = pd.Timestamp(anio, mes_numero, calendar.monthrange(anio, mes_numero)[1])

    positivas, negativas = separar_acreditaciones_negativas(liquidaciones_mes_actual)
    es_del_mes = _hasta_fecha(positivas["Fecha de Pago"], ultimo_dia_mes)

    corriente = positivas[es_del_mes][["Concepto", "Importe"]].reset_index(drop=True)
    acred_pendientes_mes_siguiente = (
        positivas[~es_del_mes][["Concepto", "Importe"]].reset_index(drop=True)
    )
    anterior = acred_mes_anterior[["Concepto", "Importe"]].reset_index(drop=True)

    esperado = round(corriente["Importe"].sum() + anterior["Importe"].sum(), 2)
    diferencia = round(total_mayor - esperado, 2)

    acreditacion_faltante = pd.DataFrame(columns=["Concepto", "Importe"])
    idx_falta_ant, idx_falta_cor = [], []

    if abs(diferencia) > TOLERANCIA_ACREDITACION:
        pool = pd.concat(
            [anterior.assign(_origen="anterior"), corriente.assign(_origen="corriente")],
            ignore_index=True
        )
        registros = list(pool.itertuples())
        objetivo = round(abs(diferencia), 2)
        encontrado = None

        tam_max = len(registros)
        if tam_max > MAX_FILAS_COMBINACION_COMPLETA:
            tam_max = MAX_TAM_COMBINACION_RECORTADA

        for tam in range(1, tam_max + 1):
            candidatos = [
                c for c in itertools.combinations(registros, tam)
                if abs(round(sum(r.Importe for r in c), 2) - objetivo) <= TOLERANCIA_ACREDITACION
            ]
            if candidatos:
                if len(candidatos) > 1:
                    _avisar(advertencias,
                            f"Acreditaciones: {len(candidatos)} combinaciones de {tam} explican "
                            f"{objetivo}, se tomó la primera.")
                encontrado = candidatos[0]
                break

        if encontrado is None:
            _avisar(advertencias,
                    f"No se encontró qué liquidación explica la diferencia de acreditaciones "
                    f"de {diferencia:,.2f}.")
        else:
            acreditacion_faltante = pd.DataFrame(
                [{"Concepto": r.Concepto, "Importe": r.Importe} for r in encontrado]
            )
            idx_falta_ant = [r.Index for r in encontrado if r._origen == "anterior"]
            idx_falta_cor = [r.Index - len(anterior) for r in encontrado if r._origen == "corriente"]

    acred_match_mes_anterior = anterior.drop(index=idx_falta_ant).reset_index(drop=True)
    acred_match_mes_corriente = corriente.drop(index=idx_falta_cor).reset_index(drop=True)

    resumen = {
        'total_mayor': total_mayor,
        'esperado':    esperado,
        'diferencia':  diferencia,
        'negativas':   len(negativas),
    }

    return (acred_match_mes_anterior, acred_match_mes_corriente,
            acred_pendientes_mes_siguiente, acreditacion_faltante, resumen)


def _encabezar_bloque_con_fecha(df, anio, mes_numero):
    """Si el bloque arranca con una celda de fecha (como en el cuadro manual),
    la actualiza al último día del mes; si no, la agrega."""
    df = df.reset_index(drop=True)
    fecha = pd.Timestamp(anio, mes_numero, calendar.monthrange(anio, mes_numero)[1])
    primera = df.iloc[0, 1] if len(df) else None
    if isinstance(primera, (pd.Timestamp, dt.date, dt.datetime)):
        df.iloc[0, 1] = fecha
        return df
    return pd.concat([_bloque(df, [fecha], [None]), df], ignore_index=True)


def _fila_vacia(df, k=1):
    return pd.DataFrame([[None] * df.shape[1]] * k, columns=df.columns).astype(object)


def _bloque(df, textos, montos):
    b = pd.DataFrame([[None] * df.shape[1] for _ in textos], columns=df.columns).astype(object)
    b.iloc[:, 1] = textos
    b.iloc[:, 3] = montos
    return b


def _texto_col(df):
    return df.iloc[:, 1].apply(lambda x: "" if pd.isna(x) else str(x).strip())


def actualizar_pendientes_mes_anterior(df_pendientes_raw, facturas_pendientes,
                                       match_facturas_pendientes):
    facturas_individuales = set(
        match_facturas_pendientes.loc[match_facturas_pendientes["Nro Factura"].notna(), "Nro Factura"]
    )

    periodos_completos = set(
        match_facturas_pendientes.loc[match_facturas_pendientes["Nro Factura"].isna(), "Periodo"]
    )
    if periodos_completos:
        facturas_individuales |= set(
            facturas_pendientes.loc[facturas_pendientes["Periodo"].isin(periodos_completos), "Nro Factura"]
        )

    patron_numerico = re.compile(r'^A-\d{4}-\d{8}')
    patron_cuentas_dia = re.compile(r'^[A-Za-z]+\d+')

    def _extraer_nro_factura(texto):
        texto = str(texto).strip()
        m = patron_numerico.match(texto)
        if m:
            return m.group()
        m = patron_cuentas_dia.match(texto)
        return m.group() if m else None

    nro_factura_por_fila = _texto_col(df_pendientes_raw).apply(_extraer_nro_factura)
    filas_a_sacar = nro_factura_por_fila.isin(facturas_individuales)

    return df_pendientes_raw[~filas_a_sacar].reset_index(drop=True)


def agregar_seccion_mes(df_pendientes_actualizado, conceptos_mes_actual,
                        falta_facturas_periodo_actual, anio=None, mes_numero=None,
                        cuadro_conceptos=None, con_anio=True, advertencias=None,
                        acreditaciones_negativas=None, agrupar_conceptos=True):
    anio, mes_numero = _resolver_mes(anio, mes_numero, cuadro_conceptos, advertencias)

    encabezado_mes = f"{MESES_ES[mes_numero]} {anio}" if con_anio else MESES_ES[mes_numero]
    if _texto_col(df_pendientes_actualizado).eq(encabezado_mes).any():
        _avisar(advertencias, f"'{encabezado_mes}' ya estaba en el cuadro de pendientes, no se duplicó.")
        return df_pendientes_actualizado

    df = df_pendientes_actualizado.reset_index(drop=True)
    ultimo_dia_mes = pd.Timestamp(anio, mes_numero, calendar.monthrange(anio, mes_numero)[1])
    nombre_mes_siguiente = MESES_ES[1] if mes_numero == 12 else MESES_ES[mes_numero + 1]

    texto = _texto_col(df)

    idx_total = [i for i, t in texto.items() if t.lower().startswith("total saldo del mes")]
    pos_total = idx_total[0] if idx_total else len(df)

    patron_gfr = re.compile(r'^Gastos Faltan Registrar Acreditaciones de .+ Registradas en .+$',
                            re.IGNORECASE)
    idx_gfr = [i for i, t in texto.items() if patron_gfr.match(t) and i > pos_total]

    bloque_viejo = _fila_vacia(df, 0)
    ini_gfr = fin_gfr = None
    if idx_gfr:
        ini_gfr = idx_gfr[0]
        fin_gfr = ini_gfr + 1
        while fin_gfr < len(df) and not df.iloc[fin_gfr].isna().all():
            fin_gfr += 1
        bloque_viejo = df.iloc[ini_gfr:fin_gfr].copy()

    def _texto_factura(fila):
        periodo = fila["Periodo"]
        if pd.isna(periodo) or not periodo:
            return str(fila["Nro Factura"])
        return f'{fila["Nro Factura"]} {periodo}'

    def _combinar_y_ordenar(conceptos, facturas):
        """Con agrupar_conceptos=True se arma como el cuadro manual: una línea
        por concepto (sumado entre liquidaciones) y las facturas una por una,
        todo ordenado alfabéticamente. Si no, se intercalan en el orden del
        cuadro de conceptos (una línea por liquidación y concepto)."""
        if agrupar_conceptos:
            items = []
            if len(conceptos):
                sumas = conceptos.groupby("Concepto", sort=True)["Monto"].sum()
                items += [(c, c, round(m, 2)) for c, m in sumas.items()]
            items += [(_texto_factura(f), _texto_factura(f), f["Importe"]) for _, f in facturas.iterrows()]
            items.sort(key=lambda x: str(x[0]).casefold())
            return items
        items = [(f["Orden"], f["Concepto"], f["Monto"]) for _, f in conceptos.iterrows()]
        items += [(f["Orden"], _texto_factura(f), f["Importe"]) for _, f in facturas.iterrows()]
        items.sort(key=lambda x: x[0])
        return items

    es_del_mes_c = _hasta_fecha(conceptos_mes_actual["Fecha de Pago"], ultimo_dia_mes)
    es_del_mes_f = _hasta_fecha(falta_facturas_periodo_actual["Fecha de Pago"], ultimo_dia_mes)

    items_mes = _combinar_y_ordenar(conceptos_mes_actual[es_del_mes_c],
                                    falta_facturas_periodo_actual[es_del_mes_f])
    items_siguiente = _combinar_y_ordenar(conceptos_mes_actual[~es_del_mes_c],
                                          falta_facturas_periodo_actual[~es_del_mes_f])

    bloque_mes = _bloque(
        df,
        [encabezado_mes, "Gastos Faltan Registrar"] + [t for _, t, _ in items_mes],
        [None, None] + [-round(m, 2) for _, _, m in items_mes]
    )

    bloque_negativas = _fila_vacia(df, 0)
    if acreditaciones_negativas is not None and len(acreditaciones_negativas):
        bloque_negativas = _bloque(
            df,
            [f'Acreditaciones Negativas {encabezado_mes}'] +
            [f'Acreditaciones Negativas {fila["Concepto"]}'
             for _, fila in acreditaciones_negativas.iterrows()],
            [None] + [-round(fila["Importe"], 2)
                      for _, fila in acreditaciones_negativas.iterrows()]
        )

    bloque_siguiente = _fila_vacia(df, 0)
    if items_siguiente:
        bloque_siguiente = _bloque(
            df,
            [f"Gastos Faltan Registrar Acreditaciones de {MESES_ES[mes_numero].lower()} "
             f"Registradas en {nombre_mes_siguiente}"] + [t for _, t, _ in items_siguiente],
            [None] + [-round(m, 2) for _, _, m in items_siguiente]
        )

    antes = df.iloc[:pos_total].copy()
    while len(antes) and antes.iloc[-1].isna().all():
        antes = antes.iloc[:-1]

    partes = [antes]
    if len(bloque_viejo):
        partes += [_fila_vacia(df), bloque_viejo]
    partes += [_fila_vacia(df), bloque_mes]
    if len(bloque_negativas):
        partes += [_fila_vacia(df), bloque_negativas]
    partes += [_fila_vacia(df, 2)]

    if ini_gfr is not None:
        partes.append(df.iloc[pos_total:ini_gfr].copy())
        if len(bloque_siguiente):
            partes.append(bloque_siguiente)
        partes.append(df.iloc[fin_gfr:].copy())
    else:
        partes.append(df.iloc[pos_total:].copy())
        if len(bloque_siguiente):
            partes += [_fila_vacia(df), bloque_siguiente]

    return pd.concat(partes, ignore_index=True)


def actualizar_acreditaciones_pendientes(pendientes_actualizado2, acred_match_mes_anterior,
                                         acred_pendientes_mes_siguiente, acreditacion_faltante,
                                         anio=None, mes_numero=None, cuadro_conceptos=None,
                                         advertencias=None):
    anio, mes_numero = _resolver_mes(anio, mes_numero, cuadro_conceptos, advertencias)

    patron_header = re.compile(r'^Acreditaciones de .+ Registradas en .+$', re.IGNORECASE)
    patron_liq = re.compile(r'^Liquidaci[oó]n .+$', re.IGNORECASE)

    texto = _texto_col(pendientes_actualizado2)
    idx_header = [i for i, t in texto.items()
                  if patron_header.match(t) and not t.lower().startswith("gastos")]

    nombre_mes = MESES_ES[mes_numero]
    nombre_mes_siguiente = MESES_ES[1] if mes_numero == 12 else MESES_ES[mes_numero + 1]
    encabezado = f"Acreditaciones de {nombre_mes} Registradas en {nombre_mes_siguiente}"

    bloque_items = [(f["Concepto"], f["Importe"]) for _, f in acreditacion_faltante.iterrows()]
    bloque_items += [(f["Concepto"], f["Importe"]) for _, f in acred_pendientes_mes_siguiente.iterrows()]

    bloque = _bloque(
        pendientes_actualizado2,
        [encabezado] + [c for c, _ in bloque_items],
        [None] + [-round(m, 2) for _, m in bloque_items]
    )

    if not idx_header:
        _avisar(advertencias, "No se encontró el bloque de acreditaciones: se agregó al final.")
        return pd.concat([pendientes_actualizado2, _fila_vacia(pendientes_actualizado2), bloque],
                         ignore_index=True)

    pos = idx_header[0]
    fin = pos + 1
    while fin < len(texto) and patron_liq.match(texto[fin]):
        fin += 1

    return pd.concat([pendientes_actualizado2.iloc[:pos], bloque,
                      pendientes_actualizado2.iloc[fin:]], ignore_index=True)


def agregar_diferencia_venta(pendientes_actualizado3, total_dean, total_hio,
                             anio=None, mes_numero=None, cuadro_conceptos=None,
                             advertencias=None):
    anio, mes_numero = _resolver_mes(anio, mes_numero, cuadro_conceptos, advertencias)

    encabezado = f"Diferencia de Venta y extracto {MESES_ES[mes_numero]} {anio}"

    df = pendientes_actualizado3.reset_index(drop=True)
    texto = _texto_col(df)
    if texto.eq(encabezado).any():
        _avisar(advertencias, f"'{encabezado}' ya estaba en el cuadro de pendientes, no se duplicó.")
        return df

    textos, montos = [encabezado], [None]
    if total_dean is not None:
        textos.append("Total D&D")
        montos.append(round(total_dean, 2))
    if total_hio is not None:
        textos.append("Total HIO")
        montos.append(round(total_hio, 2))

    # La línea de centavos del mes cierra este bloque; actualizar_saldo_mayor
    # le pone el valor.
    textos.append("Diferencia de centavos")
    montos.append(0)

    idx_dv = [i for i, t in texto.items() if t.lower().startswith("diferencia de venta y extracto")]
    if not idx_dv:
        _avisar(advertencias,
                "No se encontró ningún bloque 'Diferencia de Venta y extracto': se agregó al final.")
        antes = df.copy()
        while len(antes) and antes.iloc[-1].isna().all():
            antes = antes.iloc[:-1]
        return pd.concat([antes, _fila_vacia(df), _bloque(df, textos, montos)], ignore_index=True)

    # Va a continuación del último bloque de diferencias de venta (hasta su
    # primera fila en blanco), no al final del archivo, donde puede haber
    # restos sueltos del cuadro manual.
    fin = idx_dv[-1] + 1
    while fin < len(df) and not df.iloc[fin].isna().all():
        fin += 1

    # Una 'Diferencia de centavos' suelta después del último bloque (la que
    # dejó el cierre anterior cuando no tuvo bloque propio) se reemplaza por la
    # de este mes: la cadena de totales es acumulada, así que no se pierde nada.
    resto = df.iloc[fin:].copy()
    sueltas = _texto_col(resto).str.lower().str.startswith("diferencia de centavos")
    if sueltas.any():
        _avisar(advertencias,
                f"Se reemplazó {int(sueltas.sum())} línea/s 'Diferencia de centavos' suelta/s "
                f"por la de {MESES_ES[mes_numero]}.")
        resto = resto[~sueltas]
    while len(resto) and resto.iloc[0].isna().all():
        resto = resto.iloc[1:]

    partes = [df.iloc[:fin], _fila_vacia(df), _bloque(df, textos, montos)]
    if len(resto):
        partes += [_fila_vacia(df), resto]
    return pd.concat(partes, ignore_index=True)


def actualizar_saldo_mayor(pendientes_actualizado4, recaudacion, venta_dean, venta_hio,
                           anio, mes_numero, advertencias=None):
    ultimo_dia = calendar.monthrange(anio, mes_numero)[1]
    fecha_str = f"{ultimo_dia:02d}/{mes_numero:02d}/{anio}"
    etiqueta_hio = f"Ventas {mes_numero:02d}-{anio} HIO"
    etiqueta_dean = f"Ventas D&D {mes_numero:02d}-{anio} (online TC)"
    mes_siguiente = MESES_ES[1] if mes_numero == 12 else MESES_ES[mes_numero + 1]

    df_rec = _renombrar_columnas(recaudacion.copy(), {
        'Comentario': ('comentario',),
        'Saldo':      ('saldo',),
    })
    comentario = df_rec["Comentario"].astype(str).str.strip()
    saldos = df_rec["Saldo"].dropna()
    if saldos.empty:
        raise ValueError("El mayor de recaudación no tiene ningún saldo cargado.")
    saldo_final = round(float(saldos.iloc[-1]), 2)

    df = pendientes_actualizado4.reset_index(drop=True)

    def _fila(texto_celda, monto):
        return _bloque(df, [texto_celda], [monto])

    def _pos(prefijo, ultima=False):
        t = df.iloc[:, 1].apply(lambda x: "" if pd.isna(x) else str(x).strip().lower())
        idx = [i for i, v in t.items() if v.startswith(prefijo)]
        if not idx:
            return None
        return idx[-1] if ultima else idx[0]

    def _suma(desde, hasta):
        return round(pd.to_numeric(df.iloc[desde:hasta, 3], errors="coerce").sum(), 2)

    pos_saldo = _pos("saldo según mayor contable")
    if pos_saldo is None:
        raise ValueError("No se encontró la línea 'Saldo según Mayor Contable' en el cuadro de pendientes.")
    df.iloc[pos_saldo, 1] = f"Saldo según Mayor Contable al {fecha_str}"
    df.iloc[pos_saldo, 3] = saldo_final

    faltantes = []
    if not comentario.eq(etiqueta_hio).any():
        faltantes.append((etiqueta_hio, round(venta_hio, 2)))
    if not comentario.eq(etiqueta_dean).any():
        faltantes.append((etiqueta_dean, round(venta_dean, 2)))

    if faltantes:
        bloque = pd.concat([_fila(t, m) for t, m in faltantes], ignore_index=True)
        df = pd.concat([df.iloc[:pos_saldo + 1], bloque, df.iloc[pos_saldo + 1:]], ignore_index=True)
        _avisar(advertencias,
                f"Las ventas del mes no estaban cargadas en el mayor: se agregaron "
                f"{len(faltantes)} línea/s al cuadro.")

    # La línea 'Diferencia de centavos' del mes ya viene en el bloque de
    # 'Diferencia de Venta y extracto' que agrega agregar_diferencia_venta.

    resumen = {'saldo_mayor': saldo_final}

    pos_saldo = _pos("saldo según mayor contable")
    pos_total = _pos("total saldo del mes")
    pos_dif = _pos("diferencia para llegar al saldo")

    if pos_total is not None:
        df.iloc[pos_total, 3] = _suma(pos_saldo, pos_total)
        resumen['total_saldo_mes'] = df.iloc[pos_total, 3]

    pos_acred = _pos("acreditaciones de")
    pos_sa = _pos("saldo acreditar")
    if pos_sa is not None and pos_acred is not None and pos_dif is not None:
        monto_acred = -_suma(pos_acred, pos_dif)
        monto_fmt = f"{monto_acred:,.2f}".replace(",", "@").replace(".", ",").replace("@", ".")
        df.iloc[pos_sa, 1] = f"Saldo Acreditar el próximo mes {mes_siguiente} ({monto_fmt} )"
        resumen['saldo_acreditar'] = monto_acred

    if pos_dif is not None and pos_total is not None:
        df.iloc[pos_dif, 3] = _suma(pos_total, pos_dif)
        resumen['diferencia_saldo'] = df.iloc[pos_dif, 3]

        pos_cent = _pos("diferencia de centavos", ultima=True)
        if pos_cent is not None and pos_cent > pos_dif:
            df.iloc[pos_cent, 3] = round(-_suma(pos_dif, pos_cent), 2)
            resumen['diferencia_centavos'] = df.iloc[pos_cent, 3]

    return df, resumen


AMARILLO = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")
NEGRITA = Font(bold=True)

_MESES_TITULO = ("enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
                 "agosto", "septiembre", "octubre", "noviembre", "diciembre")
_PATRON_MES_TITULO = re.compile(rf"^({'|'.join(_MESES_TITULO)})(\s+\d{{4}})?$", re.IGNORECASE)


def _es_titulo(texto):
    t = str(texto).strip().lower()
    if not t:
        return False
    if _PATRON_MES_TITULO.match(t):
        return True
    return t.startswith((
        "gastos faltan registrar",
        "acreditaciones de",
        "total saldo del mes",
        "saldo acreditar",
        "diferencia de venta y extracto",
    ))


def _nombre_hoja_valido(nombre, advertencias=None):
    limpio = re.sub(r'[:\\/?*\[\]]', '-', nombre).strip()
    if len(limpio) > 31:
        limpio = limpio[:31].strip()
    if limpio != nombre:
        _avisar(advertencias, f"Hoja renombrada: '{nombre}' -> '{limpio}'")
    return limpio


def construir_conciliacion_xlsx(conciliacion, recaudacion, df_dean_dep, df_hio,
                                df_venta_rappi_dep, match_rappi_dean1, falta_rappi_dean1,
                                falta_dean1, match_hio_rappi1, match_hio2, match_hio_rappi2,
                                falta_rappi_hio2, falta_hio2, advertencias=None):
    hojas = [
        ("Conciliación", conciliacion),
        ("Mayor", recaudacion),
        ("Dean", df_dean_dep),
        ("HIO Combinado", df_hio),
        ("Rappi", df_venta_rappi_dep),
        ("Match Dean", match_rappi_dean1),
        ("Falta Rappi (Sobra Dean)", falta_rappi_dean1),
        ("Falta Dean (sobra Rappi)", falta_dean1),
        ("Match HIO (Localizador)", match_hio_rappi1),
        ("Match HIO - HIO (sin Localiz)", match_hio2),
        ("Match HIO - Rappi (sin Localiz)", match_hio_rappi2),
        ("Falta Rappi (sobra HIO)", falta_rappi_hio2),
        ("Falta HIO (sobra Rappi)", falta_hio2),
    ]
    con_encabezado = {nombre for nombre, _ in hojas if nombre != "Conciliación"}

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        nombres_usados = {}
        for nombre, df in hojas:
            nombre_final = _nombre_hoja_valido(nombre, advertencias)
            df.to_excel(writer, sheet_name=nombre_final, index=False,
                        header=nombre in con_encabezado)
            nombres_usados[nombre] = nombre_final

        for nombre in con_encabezado:
            ws = writer.sheets[nombres_usados[nombre]]
            for celda in ws[1]:
                celda.font = NEGRITA
            ws.freeze_panes = "A2"

        ws = writer.sheets[nombres_usados["Conciliación"]]
        for fila in range(1, ws.max_row + 1):
            texto = ws.cell(fila, 2).value
            if texto is not None and _es_titulo(texto):
                ws.cell(fila, 2).font = NEGRITA
                ws.cell(fila, 4).font = NEGRITA
                if str(texto).strip().lower().startswith("total saldo del mes"):
                    ws.cell(fila, 2).fill = AMARILLO
                    ws.cell(fila, 4).fill = AMARILLO

            celda = ws.cell(fila, 4)
            if isinstance(celda.value, (int, float)):
                celda.number_format = "#,##0.00"

        ws.column_dimensions["B"].width = 80
        ws.column_dimensions["D"].width = 18

    buf.seek(0)
    return buf


def correr_conciliacion_rappi_easa(archivos_liq, archivos_pdf, archivo_hio_documento,
                                   archivo_hio_metodo, archivo_dean, archivo_recaudacion,
                                   archivo_pendientes):
    advertencias = []

    liq_datos = [(getattr(a, 'name', f'liquidacion_{i}.xlsx'), _leer_bytes(a))
                 for i, a in enumerate(archivos_liq)]
    pdf_datos = [(getattr(a, 'name', f'factura_{i}.pdf'), _leer_bytes(a))
                 for i, a in enumerate(archivos_pdf)]

    piezas = procesar_rappi(
        [_como_archivo(datos, nombre) for nombre, datos in liq_datos],
        [_como_archivo(datos, nombre) for nombre, datos in pdf_datos],
    )
    advertencias.extend(piezas['advertencias'])
    cuadro_conceptos = cuadro_conceptos_a_df(piezas['cuadro_filas'])

    anio, mes = detectar_mes_cuadro_conceptos(cuadro_conceptos, advertencias)
    if anio is None:
        raise ValueError("No se pudo detectar el mes: el cuadro de conceptos no tiene liquidaciones.")

    df_venta_rappi = importar_ventas_rappi(liq_datos)
    if df_venta_rappi.empty:
        raise ValueError("Las liquidaciones no traen órdenes en la hoja 'Detalle'.")
    df_venta_rappi_dep = depurar_venta_rappi(df_venta_rappi)

    df_dean_rappi = _leer_reporte(archivo_dean, 'Reporte Dean',
                                  [('fecha',), ('transacc',), ('online',), ('efectivo',)])
    df_dean_dep = depurar_dean_rappi(df_dean_rappi, advertencias)

    df_hio_documento = _leer_reporte(archivo_hio_documento, 'Reporte HIO por documento',
                                     [('serie',), ('localizador',), ('venta',)])
    df_hio_metodo = _leer_reporte(archivo_hio_metodo, 'Reporte HIO por método de pago',
                                  [('serie',), ('medio', 'pago')])
    df_hio, falta_documento, falta_metodo = cruzar_hio_documento_metodo(df_hio_documento, df_hio_metodo)

    match_dean1, match_rappi_dean1, falta_rappi_dean1, falta_dean1 = cruzar_dean(
        df_venta_rappi_dep, df_dean_dep)
    match_hio, match_hio_rappi1, falta_rappi_hio1, falta_hio1 = cruzar_hio(
        df_hio, df_venta_rappi_dep)
    match_hio2, match_hio_rappi2, falta_rappi_hio2, falta_hio2 = cruzar_hio_fecha_venta(
        falta_rappi_hio1, falta_hio1)

    total_dean = calcular_total_dean(df_venta_rappi_dep, df_dean_dep)
    total_hio = calcular_total_hio(match_hio_rappi1, falta_rappi_hio2, falta_hio2)
    venta_dean = calcular_venta_dean(df_dean_dep)
    venta_hio = calcular_venta_hio(df_hio)

    df_recaudacion = _leer_reporte(archivo_recaudacion, 'Cuenta recaudación Rappi (mayor)',
                                   [('comentario',), ('haber',), ('saldo',), ('sufactura',)])
    # Sólo asientos hasta el cierre del mes: el mayor suele venir con los
    # primeros días del mes siguiente.
    df_recaudacion = filtrar_mayor_hasta_fin_de_mes(df_recaudacion, anio, mes, advertencias)

    # El archivo acumula todos los meses: se trabaja sobre el último bloque y el
    # resto se conserva como histórico.
    pendientes_completo = _leer_pendientes(archivo_pendientes, 'Cuadro de pendientes del mes anterior')
    historico, pendientes_mes_anterior = separar_historico_y_ultimo_bloque(pendientes_completo, advertencias)

    facturas_pendientes = extraer_facturas_pendientes(pendientes_mes_anterior)
    (match_facturas_pendientes, falta_facturas_pendientes,
     diferencia_factura_pendientes) = cruzar_facturas_pendientes(
        df_recaudacion, facturas_pendientes, advertencias)

    facturas_periodo_actual = extraer_facturas_cuadro_conceptos(cuadro_conceptos)
    facturas_periodo_actual = depurar_facturas_no_asignadas(
        facturas_periodo_actual, facturas_pendientes, anio, mes, advertencias)
    (match_facturas_periodo_actual, falta_facturas_periodo_actual,
     diferencia_factura_periodo_actual) = cruzar_facturas_periodo_actual(
        df_recaudacion, facturas_periodo_actual, advertencias)

    conceptos_mes_actual = extraer_conceptos_cuadro_conceptos(cuadro_conceptos)

    liquidaciones_mes_actual = extraer_liquidaciones_cuadro_conceptos(cuadro_conceptos)
    _, acreditaciones_negativas = separar_acreditaciones_negativas(liquidaciones_mes_actual)

    pendientes = actualizar_pendientes_mes_anterior(
        pendientes_mes_anterior, facturas_pendientes, match_facturas_pendientes)
    pendientes = agregar_seccion_mes(
        pendientes, conceptos_mes_actual, falta_facturas_periodo_actual,
        anio=anio, mes_numero=mes, advertencias=advertencias,
        acreditaciones_negativas=acreditaciones_negativas)

    acred_mes_anterior = extraer_acreditaciones_pendientes(pendientes_mes_anterior)
    (acred_match_mes_anterior, acred_match_mes_corriente,
     acred_pendientes_mes_siguiente, acreditacion_faltante,
     resumen_acred) = cruzar_acreditaciones(
        df_recaudacion, liquidaciones_mes_actual, acred_mes_anterior, anio, mes, advertencias)

    pendientes = actualizar_acreditaciones_pendientes(
        pendientes, acred_match_mes_anterior, acred_pendientes_mes_siguiente,
        acreditacion_faltante, anio=anio, mes_numero=mes, advertencias=advertencias)
    pendientes = agregar_diferencia_venta(
        pendientes, total_dean, total_hio, anio=anio, mes_numero=mes, advertencias=advertencias)
    conciliacion, resumen_saldo = actualizar_saldo_mayor(
        pendientes, df_recaudacion, venta_dean, venta_hio, anio, mes, advertencias)

    # La fecha que encabeza el bloque pasa a ser el cierre de este mes, y el
    # bloque nuevo se escribe a continuación del histórico (como el cuadro manual).
    conciliacion = _encabezar_bloque_con_fecha(conciliacion, anio, mes)
    if len(historico):
        conciliacion = pd.concat([historico, _fila_vacia(conciliacion, 2), conciliacion], ignore_index=True)

    buf_conciliacion = construir_conciliacion_xlsx(
        conciliacion, df_recaudacion, df_dean_dep, df_hio, df_venta_rappi_dep,
        match_rappi_dean1, falta_rappi_dean1, falta_dean1, match_hio_rappi1,
        match_hio2, match_hio_rappi2, falta_rappi_hio2, falta_hio2, advertencias)

    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        escribir_salidas_rappi(zf, piezas)
        zf.writestr('Conciliacion.xlsx', buf_conciliacion.read())
    zip_buf.seek(0)

    stats = construir_stats_rappi(piezas)
    stats['advertencias'] = advertencias
    stats['conciliacion'] = {
        'mes': mes, 'anio': anio, 'nombre_mes': MESES_ES[mes],
        'ventas_rappi': len(df_venta_rappi_dep),
        'match_dean': len(match_dean1), 'falta_rappi_dean': len(falta_rappi_dean1),
        'falta_dean': len(falta_dean1), 'match_hio': len(match_hio),
        'match_hio_fecha': len(match_hio2), 'falta_rappi_hio': len(falta_rappi_hio2),
        'falta_hio': len(falta_hio2), 'hio_sin_metodo': len(falta_documento),
        'hio_sin_documento': len(falta_metodo), 'total_dean': total_dean,
        'total_hio': total_hio, 'venta_dean': venta_dean, 'venta_hio': venta_hio,
        'match_facturas': len(match_facturas_periodo_actual),
        'falta_facturas': len(falta_facturas_periodo_actual),
        'dif_facturas': len(diferencia_factura_periodo_actual),
        'match_pendientes': len(match_facturas_pendientes),
        'falta_pendientes': len(falta_facturas_pendientes),
        'dif_pendientes': len(diferencia_factura_pendientes),
        'acred_mayor': resumen_acred['total_mayor'],
        'acred_esperado': resumen_acred['esperado'],
        'acred_diferencia': resumen_acred['diferencia'],
        'acred_faltantes': len(acreditacion_faltante),
        'acred_negativas': len(acreditaciones_negativas),
        'acred_mes_siguiente': len(acred_pendientes_mes_siguiente),
        **resumen_saldo,
    }
    return zip_buf, stats
